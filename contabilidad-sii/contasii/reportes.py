"""Salida de libros y reportes: texto para terminal, CSV y Excel (si openpyxl está instalado)."""

import csv
import os

from .libro import Libro


def pesos(n: int) -> str:
    return f"{n:,.0f}".replace(",", ".")


def encabezado(lib: Libro, titulo: str) -> str:
    c = lib.contribuyente()
    return (f"{titulo}\n{c['razon_social']} — RUT {c['rut']} — Régimen {c['regimen']}\n" + "=" * 100)


def texto_diario(lib: Libro, desde, hasta) -> str:
    out = [encabezado(lib, f"LIBRO DIARIO {desde} a {hasta}")]
    td = th = 0
    for a, lineas in lib.libro_diario(desde, hasta):
        out.append(f"\nN°{a['numero']:>6}  {a['fecha']}  [{a['tipo']}/{a['origen']}]  {a['glosa']}")
        for l in lineas:
            out.append(f"        {l['cuenta']:<10} {l['nombre'][:40]:<40} {pesos(l['debe']) if l['debe'] else '':>15} "
                       f"{pesos(l['haber']) if l['haber'] else '':>15}  {l['rut_aux'] or ''}")
            td += l["debe"]
            th += l["haber"]
    out.append("\n" + f"{'TOTALES':<59}{pesos(td):>15} {pesos(th):>15}")
    return "\n".join(out)


def texto_mayor(lib: Libro, cuenta, desde, hasta) -> str:
    nombre = lib.con.execute("SELECT nombre FROM cuenta WHERE codigo=?", (cuenta,)).fetchone()
    ini, filas = lib.libro_mayor(cuenta, desde, hasta)
    out = [encabezado(lib, f"LIBRO MAYOR {cuenta} {nombre['nombre'] if nombre else '?'} — {desde} a {hasta}"),
           f"{'Saldo inicial':<70}{pesos(ini):>15}"]
    for f in filas:
        out.append(f"N°{f['numero']:>6} {f['fecha']} {f['glosa'][:40]:<40} "
                   f"{pesos(f['debe']) if f['debe'] else '':>12} {pesos(f['haber']) if f['haber'] else '':>12} "
                   f"{pesos(f['saldo']):>15}")
    return "\n".join(out)


COLS8 = ["Débitos", "Créditos", "Deudor", "Acreedor", "Activo", "Pasivo", "Pérdida", "Ganancia"]


def texto_balance(lib: Libro, desde, hasta) -> str:
    filas, tot, res = lib.balance_8_columnas(hasta, desde)
    out = [encabezado(lib, f"BALANCE TRIBUTARIO DE 8 COLUMNAS {desde} a {hasta}"),
           f"{'Cuenta':<34}" + "".join(f"{c:>13}" for c in COLS8)]
    for cod, nom, v in filas:
        out.append(f"{cod:<8}{nom[:25]:<26}" + "".join(f"{pesos(x) if x else '':>13}" for x in v))
    out.append(f"{'SUMAS':<34}" + "".join(f"{pesos(x):>13}" for x in tot))
    ajuste = [0, 0, 0, 0, 0, 0, 0, 0]
    if res >= 0:
        ajuste[5] = ajuste[6] = res
        etiqueta = "UTILIDAD DEL EJERCICIO"
    else:
        ajuste[4] = ajuste[7] = -res
        etiqueta = "PÉRDIDA DEL EJERCICIO"
    out.append(f"{etiqueta:<34}" + "".join(f"{pesos(x) if x else '':>13}" for x in ajuste))
    tf = [a + b for a, b in zip(tot, ajuste)]
    out.append(f"{'TOTALES IGUALES':<34}" + "".join(f"{pesos(x):>13}" for x in tf))
    cuadra = tf[0] == tf[1] and tf[2] == tf[3] and tf[4] == tf[5] and tf[6] == tf[7]
    out.append("\nCuadratura: " + ("OK" if cuadra else "DESCUADRADO — revisar"))
    return "\n".join(out)


def texto_f29(lib: Libro, p, codigos: dict, vencimiento=None) -> str:
    out = [encabezado(lib, f"PROPUESTA DE TRABAJO F29 — período {p.periodo}"),
           "Borrador para cotejar con la propuesta del SII. No reemplaza la declaración."]
    if vencimiento:
        out.append(f"Vencimiento estimado: {vencimiento}")
    out.append("")
    for clave, desc, monto in p.lineas:
        cod = codigos.get(clave, "")
        out.append(f"  {('[' + cod.split(' ')[0] + ']') if cod else '':<7} {desc:<62} {pesos(monto):>15}")
    if codigos.get("_fuente"):
        out.append(f"\nCódigos del formulario: {codigos['_fuente']}")
    if p.alertas:
        out.append("\nALERTAS:")
        out += [f"  • {a}" for a in p.alertas]
    if p.verificar:
        out.append("\n[VERIFICAR] Parámetros usados sin verificación directa en la fuente:")
        out += [f"  • {v}" for v in p.verificar]
    return "\n".join(out)


# ---------------------------------------------------------------- exportación
def exportar_csv(lib: Libro, carpeta: str, desde: str, hasta: str) -> list:
    os.makedirs(carpeta, exist_ok=True)
    rutas = []
    r = os.path.join(carpeta, "libro_diario.csv")
    with open(r, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["N° asiento", "Fecha", "Tipo", "Glosa", "Cuenta", "Nombre cuenta", "Debe", "Haber", "RUT auxiliar",
                    "Hash"])
        for a, lineas in lib.libro_diario(desde, hasta):
            for l in lineas:
                w.writerow([a["numero"], a["fecha"], a["tipo"], a["glosa"], l["cuenta"], l["nombre"], l["debe"],
                            l["haber"], l["rut_aux"] or "", a["hash"]])
    rutas.append(r)
    r = os.path.join(carpeta, "balance_8_columnas.csv")
    filas, tot, _ = lib.balance_8_columnas(hasta, desde)
    with open(r, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Cuenta", "Nombre"] + COLS8)
        for cod, nom, v in filas:
            w.writerow([cod, nom] + v)
        w.writerow(["", "SUMAS"] + tot)
    rutas.append(r)
    return rutas


def exportar_excel(lib: Libro, ruta: str, desde: str, hasta: str, propuesta=None, codigos=None) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as e:
        raise RuntimeError("Instale openpyxl para exportar a Excel (pip install openpyxl) o use exportar-csv.") from e
    wb = Workbook()
    neg = Font(bold=True)
    c = lib.contribuyente()

    def hoja(titulo, cab):
        ws = wb.create_sheet(titulo)
        ws.append([f"{c['razon_social']} — RUT {c['rut']} — {titulo} — {desde} a {hasta}"])
        ws.append(cab)
        for cell in ws[2]:
            cell.font = neg
        return ws

    wb.remove(wb.active)
    ws = hoja("Diario", ["N°", "Fecha", "Tipo", "Glosa", "Cuenta", "Nombre", "Debe", "Haber", "RUT aux."])
    for a, lineas in lib.libro_diario(desde, hasta):
        for l in lineas:
            ws.append([a["numero"], a["fecha"], a["tipo"], a["glosa"], l["cuenta"], l["nombre"], l["debe"],
                       l["haber"], l["rut_aux"] or ""])

    ws = hoja("Mayor", ["Cuenta", "N°", "Fecha", "Glosa", "Debe", "Haber", "Saldo"])
    for s in lib.saldos(hasta, desde):
        ini, filas = lib.libro_mayor(s["codigo"], desde, hasta)
        ws.append([f"{s['codigo']} {s['nombre']}", "", "", "Saldo inicial", "", "", ini])
        ws.cell(ws.max_row, 1).font = neg
        for f in filas:
            ws.append(["", f["numero"], f["fecha"], f["glosa"], f["debe"], f["haber"], f["saldo"]])

    ws = hoja("Balance 8 col", ["Cuenta", "Nombre"] + COLS8)
    filas, tot, _ = lib.balance_8_columnas(hasta, desde)
    for cod, nom, v in filas:
        ws.append([cod, nom] + v)
    ws.append(["", "SUMAS"] + tot)

    for libro_doc, titulo in (("compra", "Compras"), ("venta", "Ventas"),
                              ("honorario_recibido", "Honorarios recibidos"), ("honorario_emitido", "Honorarios emitidos")):
        ws = hoja(titulo, ["Período", "Tipo DTE", "RUT", "Razón social", "Folio", "Fecha", "Exento", "Neto", "IVA",
                           "IVA no rec.", "IVA uso común", "Bruto", "Retención", "Total"])
        for d in lib.con.execute("SELECT * FROM documento WHERE libro=? AND fecha BETWEEN ? AND ? ORDER BY fecha, folio",
                                 (libro_doc, desde, hasta)):
            ws.append([d["periodo"], d["tipo_dte"], d["rut_contraparte"], d["razon_social"], d["folio"], d["fecha"],
                       d["exento"], d["neto"], d["iva"], d["iva_no_rec"], d["iva_uso_comun"], d["bruto"],
                       d["retencion"], d["total"]])

    if propuesta is not None:
        ws = hoja(f"F29 {propuesta.periodo}", ["Código", "Concepto", "Monto"])
        for clave, desc, monto in propuesta.lineas:
            ws.append([(codigos or {}).get(clave, ""), desc, monto])
        ws.append([])
        for a in propuesta.alertas:
            ws.append(["ALERTA", a])
        for v in propuesta.verificar:
            ws.append(["[VERIFICAR]", v])

    ws = wb.create_sheet("Integridad")
    problemas = lib.verificar_integridad()
    ws.append(["Verificación de cadena de hashes y cuadratura"])
    ws.append(["OK: sin problemas"] if not problemas else ["PROBLEMAS:"])
    for p in problemas:
        ws.append([p])

    for w in wb.worksheets:
        for col in w.columns:
            largo = max((len(str(x.value)) for x in col[1:] if x.value is not None), default=8)
            w.column_dimensions[col[0].column_letter].width = min(max(largo + 2, 8), 50)
    wb.save(ruta)
    return ruta
